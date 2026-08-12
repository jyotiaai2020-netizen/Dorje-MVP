from pathlib import Path
from typing import Any, Dict, Optional
import base64
import json
import subprocess
import uuid
from tempfile import TemporaryDirectory
from datetime import datetime, timezone

from fastapi import HTTPException, UploadFile, status
from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook
import imageio_ffmpeg

from app.core.config import settings


class DocumentService:
    def __init__(self, base_dir: Optional[Path] = None) -> None:
        self.base_dir = base_dir or Path(settings.UPLOADS_DIR or Path(settings.RUNTIME_ROOT) / "uploads") / "dorje-ai"
        self.base_dir.mkdir(parents=True, exist_ok=True)

    def save_upload(self, upload_file: UploadFile, conversation_id: str = "default", chat_title: str = "") -> Dict[str, Any]:
        chat_folder = self._chat_folder(conversation_id)
        attachments_dir = chat_folder / "attachments"
        attachments_dir.mkdir(parents=True, exist_ok=True)
        for folder in ("generated-text", "tables", "images", "exports", "versions"):
            (chat_folder / folder).mkdir(exist_ok=True)
        filename = self._safe_filename(upload_file.filename or "upload")
        destination = attachments_dir / filename
        size = 0
        is_video = destination.suffix.lower() == ".mp4" or upload_file.content_type == "video/mp4"
        max_upload_mb = settings.MAX_VIDEO_UPLOAD_MB if is_video else settings.MAX_UPLOAD_MB
        max_bytes = max_upload_mb * 1024 * 1024

        with destination.open("wb") as handle:
            while chunk := upload_file.file.read(1024 * 1024):
                size += len(chunk)
                if size > max_bytes:
                    handle.close()
                    destination.unlink(missing_ok=True)
                    raise HTTPException(
                        status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
                        detail=f"Files must be {max_upload_mb} MB or smaller",
                    )
                handle.write(chunk)

        content_type = upload_file.content_type or "application/octet-stream"
        content = self.extract_text(destination, content_type)
        self._update_metadata(chat_folder, filename, content_type, size, chat_title)
        return {
            "filename": filename,
            "content_type": content_type,
            "content": content,
            "characters_extracted": len(content),
            "message": "File uploaded and processed successfully",
            "chat_folder": chat_folder.name,
        }

    def extract_text(self, file_path: Path, content_type: str) -> str:
        suffix = file_path.suffix.lower()
        if content_type == "video/mp4" or suffix == ".mp4":
            return self._extract_video_frames(file_path)

        if content_type in {"image/png", "image/jpeg", "image/webp"}:
            encoded = base64.b64encode(file_path.read_bytes()).decode("ascii")
            return f"data:{content_type};base64,{encoded}"

        if content_type == "application/pdf" or suffix == ".pdf":
            try:
                reader = PdfReader(file_path)
                content = "\n\n".join(
                    text for page in reader.pages if (text := page.extract_text())
                )[:50_000]
                if not content.strip():
                    raise ValueError("No extractable PDF text")
                return content
            except Exception as exc:
                raise HTTPException(
                    status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                    detail="The PDF could not be read or contains no extractable text",
                ) from exc

        if suffix == ".docx" or content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            try:
                document = Document(file_path)
                content = "\n".join(
                    paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()
                )[:50_000]
                if not content.strip():
                    raise ValueError("No extractable DOCX text")
                return content
            except Exception as exc:
                raise HTTPException(
                    status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                    detail="The DOCX file could not be read or contains no extractable text",
                ) from exc

        if suffix == ".xlsx" or content_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
            try:
                workbook = load_workbook(file_path, read_only=True, data_only=True)
                sections = []
                for sheet in workbook.worksheets[:10]:
                    rows = ["\t".join("" if value is None else str(value) for value in row) for row in sheet.iter_rows(values_only=True)]
                    sections.append(f"SHEET: {sheet.title}\n" + "\n".join(rows[:500]))
                content = "\n\n".join(sections)[:50_000]
                if not content.strip():
                    raise ValueError("No spreadsheet content")
                return content
            except Exception as exc:
                raise HTTPException(
                    status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                    detail="The XLSX file could not be read or contains no usable cells",
                ) from exc

        if content_type.startswith("text/") or suffix in {
            ".txt", ".md", ".json", ".csv", ".py", ".js", ".ts", ".tsx"
        }:
            try:
                return file_path.read_text(encoding="utf-8")[:50_000]
            except UnicodeDecodeError:
                return file_path.read_text(encoding="latin-1")[:50_000]
        return ""

    @staticmethod
    def _extract_video_frames(file_path: Path) -> str:
        try:
            with TemporaryDirectory() as temporary_directory:
                output_pattern = str(Path(temporary_directory) / "frame-%02d.jpg")
                subprocess.run(
                    [
                        imageio_ffmpeg.get_ffmpeg_exe(), "-hide_banner", "-loglevel", "error",
                        "-i", str(file_path), "-vf", "fps=1/2,scale=512:-2", "-frames:v", "4",
                        output_pattern,
                    ],
                    check=True,
                    timeout=90,
                )
                frame_paths = sorted(Path(temporary_directory).glob("frame-*.jpg"))
                if not frame_paths:
                    raise ValueError("No video frames extracted")
                frames = [
                    "data:image/jpeg;base64," + base64.b64encode(frame.read_bytes()).decode("ascii")
                    for frame in frame_paths
                ]
                return json.dumps({"kind": "video_keyframes", "frames": frames})
        except Exception as exc:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail="The MP4 could not be read or no usable frames could be extracted",
            ) from exc

    @staticmethod
    def _safe_filename(filename: str) -> str:
        name = Path(filename).name
        return name.replace("/", "_").replace("\\", "_")

    def _chat_folder(self, conversation_id: str) -> Path:
        safe_parts = ["".join(character for character in part if character.isalnum() or character in "-_")[:80] for part in conversation_id.split("/")]
        safe_parts = [part for part in safe_parts if part][:2] or ["default"]
        return self.base_dir.joinpath(*safe_parts)

    def create_word_document(
        self,
        *,
        user_id: int,
        filename: str,
        title: str,
        body: str,
        mode: str = "create",
        document_id: str | None = None,
    ) -> Dict[str, Any]:
        """Create, rewrite, or append to a Word-compatible DOCX file.

        The file stays in the authenticated user's generated document folder and
        is also returned as base64 so the email composer can attach the exact
        same file without exposing local server paths to the browser.
        """
        safe_filename = self._safe_filename(filename or "dorje-ai-document.docx")
        if not safe_filename.lower().endswith(".docx"):
            safe_filename = f"{Path(safe_filename).stem or 'dorje-ai-document'}.docx"
        safe_document_id = "".join(character for character in (document_id or "") if character.isalnum() or character in "-_")[:80]
        if mode in {"rewrite", "append"} and safe_document_id:
            doc_id = safe_document_id
        else:
            doc_id = f"doc-{uuid.uuid4().hex[:12]}"

        documents_dir = self.base_dir / "generated-documents" / str(user_id)
        documents_dir.mkdir(parents=True, exist_ok=True)
        destination = documents_dir / f"{doc_id}.docx"

        if mode == "append" and destination.exists():
            document = Document(destination)
            if document.paragraphs:
                document.add_paragraph("")
            document.add_heading(title or "Update", level=2)
        else:
            document = Document()
            document.add_heading(title or "Dorje AI Document", level=1)

        for block in body.replace("\r\n", "\n").split("\n"):
            text = block.strip()
            if not text:
                document.add_paragraph("")
            elif text.startswith("# "):
                document.add_heading(text.removeprefix("# ").strip(), level=1)
            elif text.startswith("## "):
                document.add_heading(text.removeprefix("## ").strip(), level=2)
            elif text.startswith(("- ", "* ")):
                document.add_paragraph(text[2:].strip(), style="List Bullet")
            else:
                document.add_paragraph(text)

        document.save(destination)
        data = destination.read_bytes()
        now = datetime.now(timezone.utc).isoformat()
        metadata_path = documents_dir / f"{doc_id}.json"
        metadata_path.write_text(json.dumps({
            "document_id": doc_id,
            "filename": safe_filename,
            "title": title,
            "mode": mode,
            "updated_at": now,
            "size": len(data),
        }, indent=2), encoding="utf-8")
        return {
            "document_id": doc_id,
            "filename": safe_filename,
            "content_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "data_base64": base64.b64encode(data).decode("ascii"),
            "size": len(data),
            "text_preview": body[:600],
        }

    @staticmethod
    def _update_metadata(chat_folder: Path, filename: str, content_type: str, size: int, chat_title: str = "") -> None:
        metadata_path = chat_folder / "metadata.json"
        now = datetime.now(timezone.utc).isoformat()
        try:
            metadata = json.loads(metadata_path.read_text(encoding="utf-8")) if metadata_path.exists() else {}
        except (json.JSONDecodeError, OSError):
            metadata = {}
        metadata["chat_title"] = chat_title.strip()[:120] or metadata.get("chat_title") or chat_folder.name
        metadata.setdefault("created_date", now)
        metadata["updated_date"] = now
        metadata.setdefault("linked_files", []).append({"name": filename, "type": content_type, "size": size, "added_at": now})
        metadata.setdefault("user_edits", [])
        metadata.setdefault("generated_outputs", [])
        metadata.setdefault("export_history", [])
        metadata_path.write_text(json.dumps(metadata, indent=2), encoding="utf-8")
