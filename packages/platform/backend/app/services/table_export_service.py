from datetime import datetime
from io import BytesIO
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.schemas.dorje_ai import StructuredTable


def safe_sheet_name(title: str) -> str:
    cleaned = re.sub(r"[\\/*?:\[\]]", "-", title).strip()[:31]
    return cleaned or "DorjeAI Data"


def typed_value(value: str):
    text = str(value).strip()
    if text.startswith("="):
        return text
    normalized = text.replace(",", "").replace("$", "").replace("%", "")
    try:
        number = float(normalized)
        return int(number) if number.is_integer() else number
    except ValueError:
        return text


class TableExportService:
    def export_xlsx(self, table: StructuredTable) -> bytes:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = safe_sheet_name(table.title)
        worksheet.append(table.columns)
        border = Border(*( [Side(style="thin", color="CBD5E1")] * 4 ))
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="059669")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
        for row in table.rows:
            values = [typed_value(value) for value in (row + [""] * len(table.columns))[:len(table.columns)]]
            worksheet.append(values)
        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = border
                if isinstance(cell.value, float):
                    cell.number_format = "#,##0.00"
                elif isinstance(cell.value, int):
                    cell.number_format = "#,##0"
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        worksheet.sheet_view.showGridLines = False
        for index, column in enumerate(table.columns, start=1):
            values = [str(column), *(str(row[index - 1]) for row in table.rows if len(row) >= index)]
            worksheet.column_dimensions[get_column_letter(index)].width = min(max(max(map(len, values)) + 2, 12), 60)
        metadata = workbook.create_sheet("Metadata")
        metadata.append(["Property", "Value"])
        metadata.append(["Title", table.title])
        metadata.append(["Model", table.model_used])
        metadata.append(["Generated", table.generated_at])
        metadata.append(["Dataset ID", table.dataset_id])
        metadata.append(["Source columns", ", ".join(table.source_columns or table.columns)])
        metadata.append(["Aggregation", table.aggregation_method])
        metadata.append(["Filters", ", ".join(table.filters_applied) or "None"])
        metadata.append(["Exported", datetime.now().isoformat()])
        metadata.append(["Explanation", table.explanation])
        metadata.column_dimensions["A"].width = 18
        metadata.column_dimensions["B"].width = 80
        output = BytesIO(); workbook.save(output); output.seek(0); return output.getvalue()

    def export_pdf(self, table: StructuredTable) -> bytes:
        output = BytesIO(); styles = getSampleStyleSheet()
        body = ParagraphStyle("TableBody", parent=styles["BodyText"], fontSize=7.5, leading=9.5)
        header = ParagraphStyle("TableHeader", parent=body, textColor=colors.white, fontName="Helvetica-Bold")
        data = [[Paragraph(str(cell), header) for cell in table.columns]]
        data.extend([[Paragraph(str((row + [""] * len(table.columns))[index]), body) for index in range(len(table.columns))] for row in table.rows])
        available_width = landscape(letter)[0] - 0.8 * inch
        widths = [available_width / max(len(table.columns), 1)] * len(table.columns)
        rendered = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
        rendered.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#059669")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#94A3B8")), ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ]))
        story = [Paragraph(table.title, styles["Title"])]
        if table.explanation: story.extend([Paragraph(table.explanation, styles["BodyText"]), Spacer(1, 0.12 * inch)])
        story.extend([rendered, Spacer(1, 0.12 * inch), Paragraph(f"Generated by {table.model_used} · {table.generated_at}", styles["Italic"] )])
        document = SimpleDocTemplate(output, pagesize=landscape(letter), leftMargin=0.4 * inch, rightMargin=0.4 * inch, topMargin=0.45 * inch, bottomMargin=0.45 * inch)
        document.build(story); output.seek(0); return output.getvalue()

    def export_png(self, table: StructuredTable) -> bytes:
        font = ImageFont.load_default(); scale = 2; padding = 12 * scale; row_height = 30 * scale
        columns = max(len(table.columns), 1); column_width = 180 * scale
        width = min(columns * column_width + padding * 2, 6000)
        height = min((len(table.rows) + 2) * row_height + padding * 2, 10000)
        image = Image.new("RGB", (width, height), "white"); draw = ImageDraw.Draw(image)
        draw.text((padding, padding // 2), table.title, fill="#0F172A", font=font)
        start_y = padding + row_height
        rows = [table.columns, *table.rows]
        for row_index, row in enumerate(rows):
            y = start_y + row_index * row_height
            for column_index in range(columns):
                x = padding + column_index * column_width
                fill = "#059669" if row_index == 0 else ("#F8FAFC" if row_index % 2 == 0 else "#FFFFFF")
                draw.rectangle((x, y, x + column_width, y + row_height), fill=fill, outline="#94A3B8", width=2)
                value = str(row[column_index] if column_index < len(row) else "")
                draw.text((x + 8, y + 9), value[:60], fill="white" if row_index == 0 else "#0F172A", font=font)
        output = BytesIO(); image.save(output, format="PNG", optimize=True); output.seek(0); return output.getvalue()
